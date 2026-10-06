"""Isolated local capability probes. Only --download-models permits model downloads."""
from __future__ import annotations
import argparse
import importlib.metadata as metadata
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import urllib.request
from urllib.parse import urlparse
import wave


def local_url(value):
    parsed = urlparse(value)
    if (parsed.scheme != 'http' or parsed.hostname not in {'localhost', '127.0.0.1', '::1'}
            or parsed.username or parsed.password or parsed.query or parsed.fragment):
        raise ValueError('Service URL must be plain HTTP loopback without credentials/query')
    return value.rstrip('/')


def request(url, data=None, headers=None, timeout=10):
    # Do not send localhost through a proxy or follow a redirect to a remote host.
    class NoRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *args, **kwargs):
            return None
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}), NoRedirect())
    return json.loads(opener.open(urllib.request.Request(url, data=data, headers=headers or {}),
                                  timeout=timeout).read())


def silent_wav():
    buf = io.BytesIO()
    with wave.open(buf, 'wb') as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(16000)
        wav.writeframes(b'\0\0' * 16000)
    return buf.getvalue()


def models_in(directory):
    found = {}
    for kind in ('det', 'cls', 'rec'):
        candidates = [p for p in directory.glob('*.onnx')
                      if kind in p.name.lower() and p.stat().st_size > 0]
        if len(candidates) != 1:
            return {}
        found[kind] = str(candidates[0])
    return found


def probe(args):
    component = args.component
    if component == 'ffmpeg':
        exe = shutil.which('ffmpeg')
        if not exe:
            import imageio_ffmpeg
            exe = imageio_ffmpeg.get_ffmpeg_exe()
        subprocess.run([exe, '-version'], check=True, capture_output=True, timeout=15)
        if args.verify:
            # Synthetic tone -> actual audio decoder. No user media or output file.
            subprocess.run([exe, '-v', 'error', '-f', 'lavfi', '-i',
                            'sine=frequency=440:duration=0.1', '-f', 'null', '-'],
                           check=True, capture_output=True, timeout=15)
        return {'status': 'ready', 'executable': exe,
                'evidence': 'synthetic_decode' if args.verify else 'executable_version'}
    if component in {'ocr', 'gpu'}:
        if component == 'gpu':
            if not args.gpu_runtime or not (Path(args.gpu_runtime) / 'onnxruntime').is_dir():
                return {'status': 'missing', 'reason': 'isolated_gpu_runtime_missing'}
            sys.path.insert(0, args.gpu_runtime)
        import onnxruntime as ort
        spec = importlib.util.find_spec('rapidocr')
        directory = Path(args.ocr_models) if args.ocr_models else Path(spec.origin).parent / 'models'
        models = models_in(directory)
        if not models and not args.download_models:
            return {'status': 'missing', 'reason': 'ocr_weights_missing_or_ambiguous',
                    'model_directory': str(directory)}
        if component == 'gpu' and 'CUDAExecutionProvider' not in ort.get_available_providers():
            return {'status': 'failed', 'reason': 'cuda_provider_unavailable'}
        # Inspection never calls an engine initializer (which could fetch model weights).
        if not args.verify and not args.download_models:
            return {'status': 'detected', 'evidence': 'packages_and_weights',
                    'models': models, 'providers': ort.get_available_providers(),
                    'reason': 'run_verify_to_check_actual_inference'}
        if component == 'gpu':
            ort.preload_dlls()
        from rapidocr import RapidOCR
        params = {'EngineConfig.onnxruntime.intra_op_num_threads': 4,
                  'EngineConfig.onnxruntime.inter_op_num_threads': 1}
        for kind, path in models.items():
            params[{'det': 'Det', 'cls': 'Cls', 'rec': 'Rec'}[kind] + '.model_path'] = path
        if component == 'gpu':
            params.update({'EngineConfig.onnxruntime.use_cuda': True,
                           'EngineConfig.onnxruntime.cuda_ep_cfg.use_tf32': '0',
                           'EngineConfig.onnxruntime.cuda_ep_cfg.arena_extend_strategy': 'kSameAsRequested',
                           'EngineConfig.onnxruntime.cuda_ep_cfg.cudnn_conv_algo_search': 'HEURISTIC',
                           'EngineConfig.onnxruntime.cuda_ep_cfg.cudnn_conv_use_max_workspace': '0',
                           'EngineConfig.onnxruntime.cuda_ep_cfg.gpu_mem_limit': args.gpu_mem_limit})
        engine = RapidOCR(params=params)
        providers = [getattr(engine, n).session.session.get_providers()
                     for n in ('text_det', 'text_cls', 'text_rec')]
        if component == 'gpu':
            if any(p[0] != 'CUDAExecutionProvider' for p in providers):
                return {'status': 'failed', 'reason': 'gpu_silently_fell_back', 'providers': providers}
            for name in ('text_det', 'text_cls', 'text_rec'):
                session = getattr(engine, name).session.session
                original = session.run
                options = ort.RunOptions()
                options.add_run_config_entry('memory.enable_memory_arena_shrinkage', 'gpu:0')
                def run(outputs, inputs, run_options=None, _original=original, _options=options):
                    if run_options is not None:
                        run_options.add_run_config_entry('memory.enable_memory_arena_shrinkage', 'gpu:0')
                    return _original(outputs, inputs, run_options or _options)
                session.run = run
        from PIL import Image, ImageDraw, ImageFont
        import numpy as np
        image = Image.new('RGB', (600, 160), 'white')
        font = ImageFont.truetype('DejaVuSans.ttf', 48) if os.name != 'nt' else ImageFont.truetype('arial.ttf', 48)
        ImageDraw.Draw(image).text((25, 50), 'TEST 12345', font=font, fill='black')
        result = engine(np.array(image))
        text = ' '.join(getattr(result, 'txts', None) or [])
        if '12345' not in text:
            return {'status': 'failed', 'reason': 'synthetic_ocr_expected_digits_not_recognized',
                    'providers': providers}
        return {'status': 'ready', 'evidence': 'synthetic_text_inference', 'providers': providers,
                'runtime_version': ort.__version__, 'models': models_in(directory),
                'gpu_mem_limit': args.gpu_mem_limit if component == 'gpu' else None,
                'arena_shrinkage': component == 'gpu', 'quality_scope': 'smoke_only_not_batch_accuracy'}
    if component == 'asr':
        base = local_url(args.asr_url)
        try:
            health = request(base + '/health')
        except Exception:
            health = {}
        if health.get('ready') is True:
            if args.verify:
                boundary = 'douyin-synthetic-audio-check'
                data = (f'--{boundary}\r\nContent-Disposition: form-data; name="model"\r\n\r\nsensevoice\r\n'
                        f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="check.wav"\r\n'
                        'Content-Type: audio/wav\r\n\r\n').encode() + silent_wav() + f'\r\n--{boundary}--\r\n'.encode()
                response = request(base + '/v1/audio/transcriptions', data,
                                   {'Content-Type': f'multipart/form-data; boundary={boundary}'}, timeout=120)
                if not isinstance(response.get('text'), str):
                    raise RuntimeError('ASR response did not contain text')
            return {'status': 'ready', 'evidence': 'synthetic_audio_inference' if args.verify else 'loaded_model_health',
                    'quality_scope': 'silence_smoke_not_speech_accuracy', 'service': base}
        import funasr
        root = Path(args.asr_root)
        # Compatible with ModelScope's current and older cache layout.
        weights = [p.parent for p in (root / 'models').rglob('model.pt')
                   if p.stat().st_size > 0 and 'SenseVoiceSmall' in str(p)]
        vad = [p.parent for p in (root / 'models').rglob('model.pt')
               if p.stat().st_size > 0 and 'fsmn_vad' in str(p)]
        if (not weights or not vad) and not args.download_models:
            return {'status': 'missing', 'reason': 'asr_local_weights_missing', 'service_ready': False}
        if not args.verify and not args.download_models:
            return {'status': 'detected', 'evidence': 'packages_and_weights',
                    'reason': 'service_not_running_local_inference_not_verified'}
        os.environ['MODELSCOPE_CACHE'] = str(root / 'models')
        from funasr import AutoModel
        model = AutoModel(model=str(weights[0]) if weights else 'iic/SenseVoiceSmall',
                          vad_model=str(vad[0]) if vad else 'iic/speech_fsmn_vad_zh-cn-16k-common-pytorch',
                          device='cpu', disable_update=True, trust_remote_code=False)
        with tempfile.TemporaryDirectory(prefix='douyin-asr-check-') as temp:
            wav = Path(temp) / 'check.wav'
            wav.write_bytes(silent_wav())
            rows = model.generate(input=str(wav), language='auto', use_itn=True)
            if not isinstance(rows, list):
                raise RuntimeError('ASR inference did not return a list')
        return {'status': 'ready', 'evidence': 'direct_local_synthetic_inference',
                'service_ready': False, 'quality_scope': 'silence_smoke_not_speech_accuracy'}
    if component == 'word':
        import docx
        with tempfile.TemporaryDirectory(prefix='douyin-word-check-') as temp:
            path = Path(temp) / 'check.docx'
            document = docx.Document()
            document.add_paragraph('中文环境检查 TEST')
            document.save(path)
            if docx.Document(path).paragraphs[0].text != '中文环境检查 TEST':
                raise RuntimeError('Word roundtrip failed')
        return {'status': 'ready', 'evidence': 'chinese_docx_roundtrip', 'visual_layout': 'not_tested'}
    if component == 'collect':
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            path = Path(pw.chromium.executable_path)
            if not path.is_file():
                return {'status': 'missing', 'reason': 'chromium_binary_missing'}
            if args.verify:
                browser = pw.chromium.launch(headless=True)
                try:
                    page = browser.new_page()
                    page.set_content('<title>Local capability check</title>')
                    if page.title() != 'Local capability check':
                        raise RuntimeError('Browser probe failed')
                finally:
                    browser.close()
        return {'status': 'ready', 'evidence': 'isolated_local_browser' if args.verify else 'browser_binary',
                'douyin_login': 'not_checked', 'chrome_extension': 'not_checked'}
    raise ValueError('Unknown component')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('component', choices=['ffmpeg', 'ocr', 'asr', 'word', 'collect', 'gpu'])
    parser.add_argument('--verify', action='store_true')
    parser.add_argument('--download-models', action='store_true')
    parser.add_argument('--asr-root', default='.')
    parser.add_argument('--asr-url', default='http://127.0.0.1:18080')
    parser.add_argument('--ocr-models')
    parser.add_argument('--gpu-runtime')
    parser.add_argument('--gpu-mem-limit', type=int, default=1073741824)
    args = parser.parse_args()
    try:
        result = probe(args)
    except ModuleNotFoundError as exc:
        result = {'status': 'missing', 'reason': 'package_missing', 'module': exc.name}
    except Exception as exc:
        # Raw traceback / URLs may include private paths or credentials. Keep local diagnostics terse.
        result = {'status': 'failed', 'reason': type(exc).__name__, 'diagnostic': str(exc)[-1200:]}
    print('DOUYIN_PROBE_JSON=' + json.dumps(result, ensure_ascii=False))
    return 0 if result['status'] in {'ready', 'detected'} else 2


if __name__ == '__main__':
    raise SystemExit(main())
